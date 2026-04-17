"""
Full CSV QA Runner — item24 chatbot
Reads 145 questions from CSV, sends to chatbot at item2, captures responses,
validates with Claude API. Saves progress after every row (resume-safe).

Usage:
    python run_full_qa.py
    python run_full_qa.py --validate-only   (skip browser, only validate existing answers)
    python run_full_qa.py --limit 5         (only process first N incomplete rows, for testing)
"""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace', line_buffering=True)

import csv
import re
import time
import os
import argparse
from datetime import datetime
from playwright.sync_api import sync_playwright, TimeoutError as PWTimeout
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import anthropic

# ─── Config ────────────────────────────────────────────────────────────────
INPUT_CSV        = "C:/Users/Lucas Restrepo/Downloads/item_Testing_Claude - Sheet1 (1).csv"
OUTPUT_DIR       = "C:/Users/Lucas Restrepo/chatbot_tester"
TIMESTAMP        = datetime.now().strftime('%Y%m%d_%H%M%S')
OUTPUT_CSV       = f"{OUTPUT_DIR}/results_full_qa_{TIMESTAMP}.csv"
OUTPUT_XLSX      = f"{OUTPUT_DIR}/results_full_qa_{TIMESTAMP}.xlsx"
CHATBOT_URL      = "https://app.reshapex.ai/store/item2"
RESPONSE_TIMEOUT = 120    # seconds to wait for full response
STABILITY_SECS   = 6.0    # seconds of unchanged text = response done
MIN_RESPONSE_LEN = 50     # chars: below this we wait longer before accepting

# ─── Streaming / intermediate state detection ─────────────────────────────
# Patterns that mean the bot is still loading/streaming — NOT a final answer.
# These typically appear at the END of the text while the agent is still working.
STREAMING_END_RE = re.compile(
    r'(Loading|Searching|Checking|Fetching|Looking up|Getting|Retrieving|'
    r'Let me (check|look|search|pull|find|review|open|locate|browse|verify))\s*[\w\s,]*\.\.\.\s*$',
    re.IGNORECASE | re.DOTALL
)

# Short placeholder-only responses (the very start of streaming)
PLACEHOLDER_RE = re.compile(
    r'^(Reasoning|Thinking|Searching|Loading|\.{1,5})\s*$',
    re.IGNORECASE
)

def is_still_streaming(text: str) -> bool:
    """True if the response is not yet complete (still loading/searching)."""
    if not text or not text.strip():
        return True
    stripped = text.strip()
    # Very short placeholder (e.g. just "Reasoning")
    if PLACEHOLDER_RE.match(stripped):
        return True
    # Ends with an active loading/searching indicator
    if STREAMING_END_RE.search(stripped):
        return True
    return False


# ─── Shadow DOM helpers ────────────────────────────────────────────────────

EXTRACT_LAST_BOT_JS = """() => {
    const host = document.querySelector('reshape-chat');
    if (!host || !host.shadowRoot) return '';
    const sr = host.shadowRoot;
    const botMds = Array.from(sr.querySelectorAll('[class*="self-start"] div.markdown'));
    if (botMds.length === 0) return '';
    return botMds[botMds.length - 1].innerText.trim();
}"""

COUNT_BOT_MESSAGES_JS = """() => {
    const host = document.querySelector('reshape-chat');
    if (!host || !host.shadowRoot) return 0;
    return host.shadowRoot.querySelectorAll('[class*="self-start"] div.markdown').length;
}"""




def open_chat(page) -> bool:
    """Click the chat bubble and wait for the textarea to appear."""
    try:
        bubble = page.wait_for_selector("button[class*='rounded-full']", timeout=12000)
        bubble.click()
        page.wait_for_selector("reshape-chat >> textarea", state="visible", timeout=10000)
        time.sleep(1.5)
        return True
    except PWTimeout:
        return False


def wait_for_response(page, pre_count: int) -> str:
    """Poll until bot response is fully streamed and stable for STABILITY_SECS."""
    start        = time.time()
    stable_text  = ""
    stable_since = None

    while time.time() - start < RESPONSE_TIMEOUT:
        time.sleep(1.5)

        count = page.evaluate(COUNT_BOT_MESSAGES_JS)
        if count <= pre_count:
            continue  # bot hasn't replied yet

        current = page.evaluate(EXTRACT_LAST_BOT_JS)

        # Still streaming / loading — reset stability, keep waiting
        if is_still_streaming(current):
            stable_text  = ""
            stable_since = None
            continue

        # Text changed — reset stability timer
        if current != stable_text:
            stable_text  = current
            stable_since = None
            continue

        # Text is the same as last poll — track stability
        if stable_since is None:
            stable_since = time.time()

        elapsed_stable = time.time() - stable_since

        # Accept: stable for STABILITY_SECS AND not a stub
        if elapsed_stable >= STABILITY_SECS and len(current) >= MIN_RESPONSE_LEN:
            return current
        # Accept very short-but-stable after 2× window (e.g. "No tenemos ese producto")
        if elapsed_stable >= STABILITY_SECS * 2 and current:
            return current

    # Timeout — return whatever we managed to capture
    return stable_text or "[Timeout — no response received]"


# ─── Claude API validation ─────────────────────────────────────────────────

def validate_with_claude(client, question: str, expected: str, agent_answer: str) -> str:
    """Score agent_answer vs expected using Claude API. Returns verdict string."""
    if not agent_answer or agent_answer.startswith("[Timeout") or agent_answer.startswith("ERROR"):
        return "Incorrecto - El chatbot no respondió o hubo un error técnico al obtener la respuesta"

    expected_block = (
        expected.strip()
        if expected.strip()
        else "(Sin referencia específica — evalúa basándote en buenas prácticas de atención al cliente)"
    )

    prompt = f"""Eres un evaluador de calidad para un chatbot de atención al cliente de item24 (fabricante alemán de perfiles de aluminio y sistemas de construcción modular).

Evalúa si la respuesta del chatbot cumple con el comportamiento esperado.

PREGUNTA DEL USUARIO:
{question}

COMPORTAMIENTO ESPERADO / RESPUESTA DE REFERENCIA:
{expected_block}

RESPUESTA REAL DEL CHATBOT:
{agent_answer}

Clasifica la respuesta del chatbot en UNA de estas categorías:
- Correcto: Cubre los puntos clave del comportamiento esperado de forma precisa y completa
- Parcialmente correcto: Cubre algunos puntos pero falta información importante, hay imprecisiones, o el tono/enfoque es incorrecto
- Incorrecto: Respuesta incorrecta, engañosa, incompleta de forma grave, o falla en cubrir los aspectos clave esperados

RESPONDE EXACTAMENTE en este formato (sin texto adicional antes ni después):
[Correcto|Parcialmente correcto|Incorrecto] - [Explicación breve en 1-2 oraciones]"""

    try:
        msg = client.messages.create(
            model="claude-sonnet-4-6",
            max_tokens=400,
            messages=[{"role": "user", "content": prompt}]
        )
        return msg.content[0].text.strip()
    except Exception as e:
        return f"Validación pendiente - Error API: {str(e)[:120]}"


# ─── CSV I/O ───────────────────────────────────────────────────────────────

FIELDS = ['input', 'expected_output', 'Agent answer', 'Validation']

def load_csv(input_override: str = None) -> list:
    import glob
    if input_override:
        src = input_override
        enc = 'cp1252'
        print(f"  Starting fresh from: {os.path.basename(src)}")
    else:
        # Prefer the latest results file (has previously captured answers)
        results_files = glob.glob(os.path.join(OUTPUT_DIR, 'results_full_qa_*.csv'))
        if results_files:
            src = max(results_files, key=os.path.getmtime)
            enc = 'utf-8-sig'
            print(f"  Resuming from: {os.path.basename(src)}")
        else:
            src = INPUT_CSV
            enc = 'cp1252'
            print(f"  Starting fresh from: {os.path.basename(src)}")
    # Try cp1252 first, fall back to latin-1
    for enc_try in ([enc] if enc != 'cp1252' else ['cp1252', 'latin-1']):
        try:
            with open(src, encoding=enc_try, newline='') as f:
                rows = list(csv.DictReader(f))
            break
        except (UnicodeDecodeError, LookupError):
            continue
    return [{k.strip(): v for k, v in r.items()} for r in rows]


def save_csv(rows: list):
    with open(OUTPUT_CSV, 'w', encoding='utf-8-sig', newline='') as f:
        w = csv.DictWriter(f, fieldnames=FIELDS, extrasaction='ignore')
        w.writeheader()
        for r in rows:
            w.writerow({k: r.get(k, '') for k in FIELDS})


# ─── Excel builder ─────────────────────────────────────────────────────────

HEADER_BLUE = "1F3864"
SUB_BLUE    = "2E75B6"
GREEN       = "C6EFCE"
YELLOW      = "FFEB9C"
RED         = "FFC7CE"
ALT_ROW     = "EEF3F8"


def _fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)

def _border():
    s = Side(style="thin", color="CCCCCC")
    return Border(left=s, right=s, top=s, bottom=s)

def _verdict_color(validation: str) -> str:
    v = (validation or "").lower()
    if v.startswith("correcto") and not v.startswith("parcialmente"):
        return GREEN
    if v.startswith("parcialmente"):
        return YELLOW
    if v.startswith("incorrecto"):
        return RED
    return "F2F2F2"


def build_excel(rows: list):
    wb = openpyxl.Workbook()

    # ── Sheet 1: Resultados ─────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Resultados"

    ws.merge_cells("A1:D1")
    t = ws["A1"]
    t.value = f"item24 Chatbot QA — Resultados  |  {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    t.font = Font(bold=True, size=13, color="FFFFFF")
    t.fill = _fill(HEADER_BLUE)
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26

    col_headers = ["Pregunta (input)", "Comportamiento esperado", "Respuesta del agente", "Validación"]
    col_widths   = [55, 55, 72, 55]
    for i, (h, w) in enumerate(zip(col_headers, col_widths), 1):
        c = ws.cell(row=2, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF", size=10)
        c.fill = _fill(SUB_BLUE)
        c.alignment = Alignment(horizontal="center", wrap_text=True)
        c.border = _border()
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[2].height = 22

    for row_num, r in enumerate(rows, 3):
        alt    = ((row_num - 3) % 2 == 1)
        base   = ALT_ROW if alt else "FFFFFF"
        v_col  = _verdict_color(r.get('Validation', ''))
        vals   = [r.get('input', ''), r.get('expected_output', ''),
                  r.get('Agent answer', ''), r.get('Validation', '')]
        for col, val in enumerate(vals, 1):
            c = ws.cell(row=row_num, column=col, value=val)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = _border()
            c.font   = Font(size=9)
            c.fill   = _fill(v_col if col == 4 else base)
        ws.row_dimensions[row_num].height = 80

    ws.freeze_panes = "A3"

    # ── Sheet 2: Resumen ────────────────────────────────────────────────────
    ws2 = wb.create_sheet("Resumen")
    ws2.merge_cells("A1:C1")
    t2 = ws2["A1"]
    t2.value = "Resumen de Validaciones"
    t2.font = Font(bold=True, size=13, color="FFFFFF")
    t2.fill = _fill(HEADER_BLUE)
    t2.alignment = Alignment(horizontal="center")
    ws2.row_dimensions[1].height = 26

    correcto   = sum(1 for r in rows
                     if (r.get('Validation') or '').lower().startswith('correcto')
                     and not (r.get('Validation') or '').lower().startswith('parcialmente'))
    parcial    = sum(1 for r in rows
                     if (r.get('Validation') or '').lower().startswith('parcialmente'))
    incorrecto = sum(1 for r in rows
                     if (r.get('Validation') or '').lower().startswith('incorrecto'))
    sin_val    = len(rows) - correcto - parcial - incorrecto

    for i, h in enumerate(["Veredicto", "Cantidad", "%"], 1):
        c = ws2.cell(row=2, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF", size=10)
        c.fill = _fill(SUB_BLUE)
        c.alignment = Alignment(horizontal="center")
        c.border = _border()

    summary = [
        ("Correcto",              correcto,   GREEN),
        ("Parcialmente correcto", parcial,    YELLOW),
        ("Incorrecto",            incorrecto, RED),
        ("Sin validar / Error",   sin_val,    "F2F2F2"),
        ("TOTAL",                 len(rows),  SUB_BLUE),
    ]
    for row_num, (label, count, color) in enumerate(summary, 3):
        pct = f"{count / len(rows) * 100:.1f}%" if rows else "—"
        for col, val in enumerate([label, count, pct], 1):
            c = ws2.cell(row=row_num, column=col, value=val)
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = _border()
            c.font = Font(size=10, bold=(label == "TOTAL"), color="FFFFFF" if color == SUB_BLUE else "000000")
            c.fill = _fill(color)
        ws2.row_dimensions[row_num].height = 22

    for i, w in enumerate([32, 12, 10], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    wb.save(OUTPUT_XLSX)
    print(f"  Excel → {OUTPUT_XLSX}")


# ─── Main ──────────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--validate-only', action='store_true',
                        help='Skip browser; only validate pre-existing answers with Claude API')
    parser.add_argument('--limit', type=int, default=0,
                        help='Only process first N incomplete rows (useful for smoke-testing)')
    parser.add_argument('--clean-incomplete', action='store_true',
                        help='Clear Agent answer cells that contain streaming artifacts, then re-run')
    parser.add_argument('--rerun-incorrect', action='store_true',
                        help='Clear and re-run rows where Validation starts with Incorrecto')
    parser.add_argument('--input', type=str, default=None,
                        help='Path to input CSV (bypasses auto-resume; use for a new test file)')
    args = parser.parse_args()

    print(f"\n{'='*65}")
    print("  item24 Chatbot Full QA Runner")
    print(f"  URL: {CHATBOT_URL}")
    print(f"  Input: {INPUT_CSV}")
    print(f"  Output: {OUTPUT_CSV}")
    print(f"{'='*65}\n")

    rows = load_csv(input_override=args.input)
    print(f"Loaded {len(rows)} rows")

    # --clean-incomplete: clear rows with streaming artifacts
    if args.clean_incomplete:
        cleaned = 0
        for r in rows:
            ans = r.get('Agent answer', '').strip()
            if ans and is_still_streaming(ans):
                r['Agent answer'] = ''
                r['Validation']   = ''
                cleaned += 1
        print(f"  Cleaned {cleaned} streaming/incomplete responses — will re-run those")

    # --rerun-incorrect: clear rows where Validation = Incorrecto so they get re-run
    if args.rerun_incorrect:
        cleared = 0
        for r in rows:
            if (r.get('Validation') or '').lower().startswith('incorrecto'):
                r['Agent answer'] = ''
                r['Validation']   = ''
                cleared += 1
        print(f"  Cleared {cleared} 'Incorrecto' rows — will re-run those")

    need_browser    = sum(1 for r in rows if not r.get('Agent answer', '').strip())
    need_validation = sum(1 for r in rows if not r.get('Validation', '').strip())
    already_done    = sum(1 for r in rows if r.get('Agent answer', '').strip()
                                          and r.get('Validation', '').strip())
    print(f"  Complete (skip):     {already_done}")
    print(f"  Need browser:        {need_browser}")
    print(f"  Need validation:     {need_validation}")

    # Init Anthropic — check env var, then local config file, then prompt
    api_key = os.environ.get('ANTHROPIC_API_KEY', '').strip()

    if not api_key:
        # Check local key file (user can put their key there once)
        key_file = os.path.join(OUTPUT_DIR, 'anthropic_key.txt')
        if os.path.isfile(key_file):
            with open(key_file) as kf:
                api_key = kf.read().strip()
            if api_key:
                print("  Anthropic API:       key loaded from anthropic_key.txt")

    if not api_key:
        print("\n  ANTHROPIC_API_KEY not found.")
        print("  Options:")
        print("    a) Set env var:  set ANTHROPIC_API_KEY=sk-ant-...")
        print(f"   b) Create file:  {OUTPUT_DIR}\\anthropic_key.txt  (paste key inside)")
        print("    c) Enter key now (will NOT be saved)")
        try:
            entered = input("  Enter API key (or press Enter to skip validation): ").strip()
            if entered:
                api_key = entered
        except (EOFError, KeyboardInterrupt):
            pass

    if api_key:
        client = anthropic.Anthropic(api_key=api_key)
        print("  Anthropic API:       OK")
    else:
        client = None
        print("  Anthropic API:       SKIPPED — answers will be collected but not validated")

    # Save initial CSV copy
    save_csv(rows)
    print(f"\nSaved initial CSV → {OUTPUT_CSV}")

    processed = 0

    with sync_playwright() as p:
        browser = p.chromium.launch(headless=False, slow_mo=80)

        for idx, row in enumerate(rows):
            question       = (row.get('input') or '').strip()
            expected       = (row.get('expected_output') or '').strip()
            has_answer     = bool(row.get('Agent answer', '').strip())
            has_validation = bool(row.get('Validation', '').strip())

            # Skip fully complete rows
            if has_answer and has_validation:
                print(f"[{idx+1:3d}/{len(rows)}] SKIP  {question[:60]}")
                continue

            if args.limit and processed >= args.limit:
                print(f"\n--limit {args.limit} reached, stopping early.")
                break

            print(f"\n[{idx+1:3d}/{len(rows)}] {'VALIDATE-ONLY' if (has_answer or args.validate_only) else 'CHAT+VALIDATE'}")
            print(f"  Q: {question[:100]}")

            # ── Step 1: Get chatbot response (unless --validate-only or already have it)
            if not has_answer and not args.validate_only:
                page = None
                try:
                    page = browser.new_page(viewport={"width": 1280, "height": 900})
                    page.goto(CHATBOT_URL, wait_until="networkidle", timeout=35000)
                    time.sleep(2)

                    if not open_chat(page):
                        raise RuntimeError("Could not open chat bubble or textarea")

                    pre_count = page.evaluate(COUNT_BOT_MESSAGES_JS)
                    ta = page.locator("reshape-chat >> textarea").first
                    ta.click()
                    ta.fill(question)
                    time.sleep(0.3)
                    ta.press("Enter")

                    response = wait_for_response(page, pre_count)
                    row['Agent answer'] = response
                    has_answer = True
                    print(f"  → {len(response)} chars: {response[:120].replace(chr(10),' ')}...")

                except Exception as e:
                    row['Agent answer'] = f"ERROR: {str(e)[:200]}"
                    row['Validation']   = "Incorrecto - Error técnico al obtener respuesta del chatbot"
                    print(f"  ✗ Browser error: {e}")
                    save_csv(rows)
                    processed += 1
                    if page:
                        try: page.close()
                        except: pass
                    time.sleep(1)
                    continue
                finally:
                    if page:
                        try: page.close()
                        except: pass

            # ── Step 2: Validate with Claude
            if has_answer and not has_validation and client:
                verdict = validate_with_claude(client, question, expected, row.get('Agent answer', ''))
                row['Validation'] = verdict
                print(f"  → {verdict[:110]}")
            elif not client and not has_validation:
                row['Validation'] = "Validación pendiente - ANTHROPIC_API_KEY no configurada"

            # Save progress after every row
            save_csv(rows)
            processed += 1
            time.sleep(0.5)

        browser.close()

    # ── Final Excel ─────────────────────────────────────────────────────────
    print(f"\n{'='*65}")
    print("Building Excel report...")
    build_excel(rows)

    # ── Summary ─────────────────────────────────────────────────────────────
    correcto   = sum(1 for r in rows
                     if (r.get('Validation') or '').lower().startswith('correcto')
                     and not (r.get('Validation') or '').lower().startswith('parcialmente'))
    parcial    = sum(1 for r in rows
                     if (r.get('Validation') or '').lower().startswith('parcialmente'))
    incorrecto = sum(1 for r in rows
                     if (r.get('Validation') or '').lower().startswith('incorrecto'))

    print(f"\n{'='*65}")
    print(f"  RESULTADOS FINALES ({len(rows)} preguntas total)")
    print(f"  Correcto:              {correcto:3d}  ({correcto/len(rows)*100:.1f}%)")
    print(f"  Parcialmente correcto: {parcial:3d}  ({parcial/len(rows)*100:.1f}%)")
    print(f"  Incorrecto:            {incorrecto:3d}  ({incorrecto/len(rows)*100:.1f}%)")
    print(f"{'='*65}")
    print(f"  CSV   → {OUTPUT_CSV}")
    print(f"  Excel → {OUTPUT_XLSX}")
    print(f"{'='*65}\n")


if __name__ == "__main__":
    main()
